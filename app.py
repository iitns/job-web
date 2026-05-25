import json
import logging
import os
import re
from pathlib import Path
from typing import Any

import psycopg2
import psycopg2.extras
import requests
from docx import Document
from flask import Flask, jsonify, request, send_from_directory
from werkzeug.utils import secure_filename


logging.basicConfig(level=os.environ.get("LOG_LEVEL", "INFO"))
logger = logging.getLogger(__name__)

BASE_DIR = Path(__file__).resolve().parent
DIST_DIR = BASE_DIR / "dist"
RESUME_STORAGE_DIR = Path(os.environ.get("RESUME_STORAGE_DIR", "/data/resumes"))

app = Flask(__name__, static_folder=str(DIST_DIR), static_url_path="")

PG_CONFIG = {
    "host": os.environ.get("PG_HOST", "postgresql-service.storage.svc.cluster.local"),
    "port": int(os.environ.get("PG_PORT", "5432")),
    "dbname": os.environ.get("PG_DB", "job_recommender"),
    "user": os.environ.get("PG_USER", "postgres"),
    "password": os.environ.get("PG_PASSWORD", ""),
}

ES_HOST = os.environ.get("ES_HOST", "http://elasticsearch-service.default.svc.cluster.local:9200").rstrip("/")
ES_INDEX = os.environ.get("ES_INDEX", "job_recommender_jobs")
LLM_BASE_URL = os.environ.get("LLM_BASE_URL", "").rstrip("/")
LLM_API_KEY = os.environ.get("LLM_API_KEY", "")
LLM_MODEL = os.environ.get("LLM_MODEL", "")

DEFAULT_PAGE_SIZE = 12
MAX_PAGE_SIZE = 50
ALLOWED_RESUME_EXTENSIONS = {".docx"}

RESUME_TABLE = "job_web_resumes"
RESUME_PROFILE_TABLE = "job_web_resume_profiles"
RESUME_SCHEMA_PATH = BASE_DIR / "migrations" / "001_create_job_web_resume_tables.sql"


def get_page_args() -> tuple[int, int]:
    page = max(1, int(request.args.get("page", "1")))
    page_size = min(MAX_PAGE_SIZE, max(1, int(request.args.get("page_size", str(DEFAULT_PAGE_SIZE)))))
    return page, page_size


def get_companies_arg() -> list[str]:
    raw = request.args.getlist("company")
    if not raw:
        csv_value = request.args.get("companies", "")
        raw = csv_value.split(",") if csv_value else []
    return [value.strip() for value in raw if value and value.strip()]


def db_connect():
    return psycopg2.connect(**PG_CONFIG)


def ensure_runtime_state() -> None:
    RESUME_STORAGE_DIR.mkdir(parents=True, exist_ok=True)
    with db_connect() as conn:
        with conn.cursor() as cur:
            cur.execute(RESUME_SCHEMA_PATH.read_text(encoding="utf-8"))
        conn.commit()


def summarize_text(*parts: Any, limit: int = 220) -> str | None:
    text = " ".join(str(part or "").strip() for part in parts if str(part or "").strip())
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return None
    if len(text) <= limit:
        return text
    return f"{text[: limit - 1].rstrip()}…"


def company_mark(company: str | None) -> str:
    words = [chunk for chunk in re.split(r"[^A-Za-z0-9]+", str(company or "").strip()) if chunk]
    if not words:
        return "?"
    if len(words) == 1:
        return words[0][:2].upper()
    return f"{words[0][0]}{words[1][0]}".upper()


def serialize_job(row: dict[str, Any]) -> dict[str, Any]:
    summary = summarize_text(
        row.get("summary"),
        row.get("team_description"),
        row.get("responsibilities"),
        row.get("raw_description"),
    )
    return {
        "job_id": row.get("job_id"),
        "company": row.get("company"),
        "company_mark": company_mark(row.get("company")),
        "company_logo_url": row.get("company_logo_url") or row.get("logo_url"),
        "title": row.get("title"),
        "location": row.get("location"),
        "locations": row.get("locations") or [],
        "primary_location": row.get("location") or ", ".join(row.get("locations") or []),
        "level_guess": row.get("level_guess"),
        "team": row.get("team"),
        "display_team": row.get("team") or row.get("level_guess"),
        "raw_description": row.get("raw_description"),
        "team_description": row.get("team_description"),
        "summary": summary,
        "responsibilities": row.get("responsibilities"),
        "minimum_qualifications": row.get("minimum_qualifications"),
        "preferred_qualifications": row.get("preferred_qualifications"),
        "skills": row.get("skills") or [],
        "domains": row.get("domains") or [],
        "url": row.get("url"),
        "posted_at": row.get("posted_at").isoformat() if row.get("posted_at") else None,
        "first_seen_at": row.get("first_seen_at").isoformat() if row.get("first_seen_at") else None,
        "last_seen_at": row.get("last_seen_at").isoformat() if row.get("last_seen_at") else None,
        "is_active": row.get("is_active", True),
        "detail_extracted_at": row.get("detail_extracted_at").isoformat() if row.get("detail_extracted_at") else None,
    }


def serialize_resume(row: dict[str, Any] | None) -> dict[str, Any] | None:
    if not row:
        return None

    return {
        "id": row.get("id"),
        "filename": row.get("filename"),
        "content_type": row.get("content_type"),
        "storage_path": row.get("storage_path"),
        "file_size_bytes": row.get("file_size_bytes"),
        "status": row.get("status"),
        "raw_text": row.get("raw_text"),
        "extract_error": row.get("extract_error"),
        "llm_model": row.get("llm_model"),
        "normalization_method": row.get("normalization_method"),
        "normalized_at": row.get("normalized_at").isoformat() if row.get("normalized_at") else None,
        "created_at": row.get("created_at").isoformat() if row.get("created_at") else None,
        "updated_at": row.get("updated_at").isoformat() if row.get("updated_at") else None,
    }


def serialize_resume_profile(row: dict[str, Any] | None) -> dict[str, Any] | None:
    if not row:
        return None

    return {
        "resume_id": row.get("resume_id"),
        "candidate_name": row.get("candidate_name"),
        "summary": row.get("summary"),
        "years_of_experience": float(row["years_of_experience"]) if row.get("years_of_experience") is not None else None,
        "current_title": row.get("current_title"),
        "seniority": row.get("seniority"),
        "locations": row.get("locations") or [],
        "skills": row.get("skills") or [],
        "domains": row.get("domains") or [],
        "companies": row.get("companies") or [],
        "roles": row.get("roles") or [],
        "preferred_job_titles": row.get("preferred_job_titles") or [],
        "preferred_locations": row.get("preferred_locations") or [],
        "team_keywords": row.get("team_keywords") or [],
        "responsibility_keywords": row.get("responsibility_keywords") or [],
        "qualification_keywords": row.get("qualification_keywords") or [],
        "education": row.get("education") or [],
        "certifications": row.get("certifications") or [],
        "projects": row.get("projects") or [],
        "experience_items": row.get("experience_items") or [],
        "raw_profile_json": row.get("raw_profile_json") or {},
        "created_at": row.get("created_at").isoformat() if row.get("created_at") else None,
        "updated_at": row.get("updated_at").isoformat() if row.get("updated_at") else None,
    }


def fetch_company_options(conn) -> list[str]:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT DISTINCT company
            FROM jobs
            WHERE is_active = TRUE
            ORDER BY company ASC
            """
        )
        return [row[0] for row in cur.fetchall()]


def fetch_jobs_from_postgres(*, page: int, page_size: int, companies: list[str]) -> dict[str, Any]:
    filters = ["is_active = TRUE"]
    params: list[Any] = []

    if companies:
        filters.append("company = ANY(%s)")
        params.append(companies)

    where_clause = " AND ".join(filters)
    offset = (page - 1) * page_size

    with db_connect() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(f"SELECT COUNT(*) AS total FROM jobs WHERE {where_clause}", params)
            total = cur.fetchone()["total"]

            cur.execute(
                f"""
                SELECT
                    job_id,
                    company,
                    title,
                    location,
                    locations,
                    level_guess,
                    team,
                    raw_description,
                    team_description,
                    responsibilities,
                    minimum_qualifications,
                    preferred_qualifications,
                    skills,
                    domains,
                    url,
                    posted_at,
                    first_seen_at,
                    last_seen_at,
                    is_active,
                    detail_extracted_at
                FROM jobs
                WHERE {where_clause}
                ORDER BY posted_at DESC NULLS LAST, last_seen_at DESC, job_id ASC
                OFFSET %s
                LIMIT %s
                """,
                [*params, offset, page_size],
            )
            jobs = [serialize_job(dict(row)) for row in cur.fetchall()]

        companies_all = fetch_company_options(conn)

    return {
        "jobs": jobs,
        "total": total,
        "page": page,
        "page_size": page_size,
        "companies": companies_all,
        "source": "postgres",
    }


def search_jobs_in_elasticsearch(*, keyword: str, page: int, page_size: int, companies: list[str]) -> dict[str, Any]:
    from_offset = (page - 1) * page_size
    query: dict[str, Any] = {
        "from": from_offset,
        "size": page_size,
        "track_total_hits": True,
        "_source": [
            "job_id",
            "company",
            "title",
            "location",
            "locations",
            "level_guess",
            "team",
            "raw_description",
            "team_description",
            "responsibilities",
            "minimum_qualifications",
            "preferred_qualifications",
            "skills",
            "domains",
            "url",
            "posted_at",
            "first_seen_at",
            "last_seen_at",
            "is_active",
            "detail_extracted_at",
        ],
        "query": {
            "bool": {
                "must": [
                    {
                        "multi_match": {
                            "query": keyword,
                            "fields": [
                                "title^4",
                                "company^3",
                                "team^2",
                                "location^2",
                                "team_description",
                                "responsibilities",
                                "minimum_qualifications",
                                "preferred_qualifications",
                                "skills^2",
                                "domains",
                            ],
                            "type": "best_fields",
                            "operator": "and",
                        }
                    }
                ],
                "filter": [{"term": {"is_active": True}}],
            }
        },
        "sort": [
            "_score",
            {"posted_at": {"order": "desc", "missing": "_last"}},
            {"last_seen_at": {"order": "desc", "missing": "_last"}},
        ],
    }

    if companies:
        query["query"]["bool"]["filter"].append({"terms": {"company": companies}})

    response = requests.post(
        f"{ES_HOST}/{ES_INDEX}/_search",
        json=query,
        timeout=10,
    )

    if response.status_code == 404:
        logger.warning("Elasticsearch index %s not found", ES_INDEX)
        with db_connect() as conn:
            companies_all = fetch_company_options(conn)
        return {
            "jobs": [],
            "total": 0,
            "page": page,
            "page_size": page_size,
            "companies": companies_all,
            "source": "elasticsearch",
            "search_ready": False,
        }

    response.raise_for_status()
    payload = response.json()
    hits = payload.get("hits", {})
    total = hits.get("total", {}).get("value", 0)
    jobs = [serialize_job(hit.get("_source", {})) for hit in hits.get("hits", [])]

    with db_connect() as conn:
        companies_all = fetch_company_options(conn)

    return {
        "jobs": jobs,
        "total": total,
        "page": page,
        "page_size": page_size,
        "companies": companies_all,
        "source": "elasticsearch",
        "search_ready": True,
    }


def allowed_resume_file(filename: str) -> bool:
    return Path(filename).suffix.lower() in ALLOWED_RESUME_EXTENSIONS


def save_uploaded_resume(file_storage) -> tuple[Path, int]:
    safe_name = secure_filename(file_storage.filename or "resume.docx")
    target_path = RESUME_STORAGE_DIR / safe_name
    suffix = 1
    while target_path.exists():
        target_path = RESUME_STORAGE_DIR / f"{Path(safe_name).stem}-{suffix}{Path(safe_name).suffix}"
        suffix += 1
    file_storage.save(target_path)
    return target_path, target_path.stat().st_size


def extract_text_from_docx(docx_path: Path) -> str:
    document = Document(str(docx_path))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def dedupe_text_items(items: list[Any]) -> list[str]:
    seen: set[str] = set()
    normalized: list[str] = []
    for item in items:
        value = str(item or "").strip()
        if not value:
            continue
        lowered = value.lower()
        if lowered in seen:
            continue
        seen.add(lowered)
        normalized.append(value)
    return normalized


def parse_json_payload(text: str) -> dict[str, Any]:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    return json.loads(text)


def normalize_resume_with_llm(raw_text: str) -> tuple[dict[str, Any], dict[str, Any]]:
    if not (LLM_BASE_URL and LLM_API_KEY and LLM_MODEL):
        return fallback_resume_profile(raw_text), {
            "normalization_method": "fallback",
            "llm_model": None,
        }

    prompt = {
        "role": "user",
        "content": (
            "Normalize the following resume into JSON. "
            "Return only valid JSON with these keys: "
            "candidate_name, summary, years_of_experience, current_title, seniority, "
            "locations, skills, domains, companies, roles, preferred_job_titles, "
            "preferred_locations, team_keywords, responsibility_keywords, qualification_keywords, "
            "education, certifications, projects, experience_items. "
            "education, projects, and experience_items must be arrays. "
            "skills/domains/companies/roles and keyword fields must be arrays of strings. "
            "Focus on fields that map well to software engineering job postings.\n\n"
            f"Resume:\n{raw_text}"
        ),
    }

    response = requests.post(
        f"{LLM_BASE_URL}/chat/completions",
        headers={
            "Authorization": f"Bearer {LLM_API_KEY}",
            "Content-Type": "application/json",
        },
        json={
            "model": LLM_MODEL,
            "messages": [
                {
                    "role": "system",
                    "content": "You extract structured recruiting profiles from resumes. Output JSON only.",
                },
                prompt,
            ],
            "temperature": 0.2,
            "response_format": {"type": "json_object"},
        },
        timeout=60,
    )
    response.raise_for_status()
    payload = response.json()
    content = payload["choices"][0]["message"]["content"]
    parsed = parse_json_payload(content)
    return parsed, {
        "normalization_method": "llm",
        "llm_model": LLM_MODEL,
    }


def fallback_resume_profile(raw_text: str) -> dict[str, Any]:
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    summary = " ".join(lines[:5])[:800]
    candidate_name = lines[0] if lines else None

    known_skills = [
        "Python", "Java", "JavaScript", "TypeScript", "React", "Vue", "Node.js",
        "PostgreSQL", "MySQL", "MongoDB", "Redis", "Elasticsearch", "Docker",
        "Kubernetes", "AWS", "GCP", "Azure", "Flask", "Django", "FastAPI",
        "Spring", "Go", "Rust", "C++", "C#", "TensorFlow", "PyTorch",
    ]
    found_skills = [skill for skill in known_skills if re.search(rf"\b{re.escape(skill)}\b", raw_text, re.IGNORECASE)]

    companies = []
    roles = []
    for line in lines[:30]:
        if re.search(r"\b(engineer|developer|manager|lead|architect)\b", line, re.IGNORECASE):
            roles.append(line[:120])

    return {
        "candidate_name": candidate_name,
        "summary": summary,
        "years_of_experience": None,
        "current_title": roles[0] if roles else None,
        "seniority": None,
        "locations": [],
        "skills": dedupe_text_items(found_skills),
        "domains": [],
        "companies": dedupe_text_items(companies),
        "roles": dedupe_text_items(roles),
        "preferred_job_titles": [],
        "preferred_locations": [],
        "team_keywords": [],
        "responsibility_keywords": [],
        "qualification_keywords": dedupe_text_items(found_skills),
        "education": [],
        "certifications": [],
        "projects": [],
        "experience_items": [],
    }


def sanitize_profile(profile: dict[str, Any]) -> dict[str, Any]:
    years_of_experience = profile.get("years_of_experience")
    try:
        years_of_experience = float(years_of_experience) if years_of_experience is not None else None
    except (TypeError, ValueError):
        years_of_experience = None

    return {
        "candidate_name": profile.get("candidate_name"),
        "summary": profile.get("summary"),
        "years_of_experience": years_of_experience,
        "current_title": profile.get("current_title"),
        "seniority": profile.get("seniority"),
        "locations": dedupe_text_items(profile.get("locations", [])),
        "skills": dedupe_text_items(profile.get("skills", [])),
        "domains": dedupe_text_items(profile.get("domains", [])),
        "companies": dedupe_text_items(profile.get("companies", [])),
        "roles": dedupe_text_items(profile.get("roles", [])),
        "preferred_job_titles": dedupe_text_items(profile.get("preferred_job_titles", [])),
        "preferred_locations": dedupe_text_items(profile.get("preferred_locations", [])),
        "team_keywords": dedupe_text_items(profile.get("team_keywords", [])),
        "responsibility_keywords": dedupe_text_items(profile.get("responsibility_keywords", [])),
        "qualification_keywords": dedupe_text_items(profile.get("qualification_keywords", [])),
        "education": profile.get("education", []),
        "certifications": dedupe_text_items(profile.get("certifications", [])),
        "projects": profile.get("projects", []),
        "experience_items": profile.get("experience_items", []),
    }


def insert_resume_record(*, filename: str, content_type: str | None, storage_path: Path, file_size_bytes: int) -> int:
    with db_connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO job_web_resumes (filename, content_type, storage_path, file_size_bytes, status)
                VALUES (%s, %s, %s, %s, 'uploaded')
                RETURNING id
                """,
                (filename, content_type, str(storage_path), file_size_bytes),
            )
            resume_id = cur.fetchone()[0]
        conn.commit()
    return resume_id


def update_resume_text(*, resume_id: int, raw_text: str) -> None:
    with db_connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                UPDATE job_web_resumes
                SET raw_text = %s,
                    status = 'text_extracted',
                    updated_at = NOW()
                WHERE id = %s
                """,
                (raw_text, resume_id),
            )
        conn.commit()


def update_resume_failure(*, resume_id: int, message: str) -> None:
    with db_connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                UPDATE job_web_resumes
                SET status = 'failed',
                    extract_error = %s,
                    updated_at = NOW()
                WHERE id = %s
                """,
                (message, resume_id),
            )
        conn.commit()


def upsert_resume_profile(*, resume_id: int, profile: dict[str, Any], metadata: dict[str, Any]) -> None:
    profile = sanitize_profile(profile)
    with db_connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO job_web_resume_profiles (
                    resume_id, candidate_name, summary, years_of_experience, current_title,
                    seniority, locations, skills, domains, companies, roles,
                    preferred_job_titles, preferred_locations, team_keywords,
                    responsibility_keywords, qualification_keywords, education,
                    certifications, projects, experience_items, raw_profile_json
                )
                VALUES (
                    %(resume_id)s, %(candidate_name)s, %(summary)s, %(years_of_experience)s, %(current_title)s,
                    %(seniority)s, %(locations)s, %(skills)s, %(domains)s, %(companies)s, %(roles)s,
                    %(preferred_job_titles)s, %(preferred_locations)s, %(team_keywords)s,
                    %(responsibility_keywords)s, %(qualification_keywords)s, %(education)s,
                    %(certifications)s, %(projects)s, %(experience_items)s, %(raw_profile_json)s
                )
                ON CONFLICT (resume_id) DO UPDATE SET
                    candidate_name = EXCLUDED.candidate_name,
                    summary = EXCLUDED.summary,
                    years_of_experience = EXCLUDED.years_of_experience,
                    current_title = EXCLUDED.current_title,
                    seniority = EXCLUDED.seniority,
                    locations = EXCLUDED.locations,
                    skills = EXCLUDED.skills,
                    domains = EXCLUDED.domains,
                    companies = EXCLUDED.companies,
                    roles = EXCLUDED.roles,
                    preferred_job_titles = EXCLUDED.preferred_job_titles,
                    preferred_locations = EXCLUDED.preferred_locations,
                    team_keywords = EXCLUDED.team_keywords,
                    responsibility_keywords = EXCLUDED.responsibility_keywords,
                    qualification_keywords = EXCLUDED.qualification_keywords,
                    education = EXCLUDED.education,
                    certifications = EXCLUDED.certifications,
                    projects = EXCLUDED.projects,
                    experience_items = EXCLUDED.experience_items,
                    raw_profile_json = EXCLUDED.raw_profile_json,
                    updated_at = NOW()
                """,
                {
                    "resume_id": resume_id,
                    "candidate_name": profile["candidate_name"],
                    "summary": profile["summary"],
                    "years_of_experience": profile["years_of_experience"],
                    "current_title": profile["current_title"],
                    "seniority": profile["seniority"],
                    "locations": profile["locations"],
                    "skills": profile["skills"],
                    "domains": profile["domains"],
                    "companies": profile["companies"],
                    "roles": profile["roles"],
                    "preferred_job_titles": profile["preferred_job_titles"],
                    "preferred_locations": profile["preferred_locations"],
                    "team_keywords": profile["team_keywords"],
                    "responsibility_keywords": profile["responsibility_keywords"],
                    "qualification_keywords": profile["qualification_keywords"],
                    "education": psycopg2.extras.Json(profile["education"]),
                    "certifications": profile["certifications"],
                    "projects": psycopg2.extras.Json(profile["projects"]),
                    "experience_items": psycopg2.extras.Json(profile["experience_items"]),
                    "raw_profile_json": psycopg2.extras.Json(profile),
                },
            )
            cur.execute(
                """
                UPDATE job_web_resumes
                SET status = 'normalized',
                    llm_model = %s,
                    normalization_method = %s,
                    normalized_at = NOW(),
                    extract_error = NULL,
                    updated_at = NOW()
                WHERE id = %s
                """,
                (metadata.get("llm_model"), metadata.get("normalization_method"), resume_id),
            )
        conn.commit()


def fetch_resume_and_profile(*, resume_id: int | None = None, latest: bool = False) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    clause = "r.id = %s"
    params: tuple[Any, ...] = (resume_id,)
    if latest:
        clause = "TRUE"
        params = ()

    with db_connect() as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(
                f"""
                SELECT
                    r.id,
                    r.filename,
                    r.content_type,
                    r.storage_path,
                    r.file_size_bytes,
                    r.status,
                    r.raw_text,
                    r.extract_error,
                    r.llm_model,
                    r.normalization_method,
                    r.normalized_at,
                    r.created_at,
                    r.updated_at
                FROM job_web_resumes r
                WHERE {clause}
                ORDER BY r.created_at DESC
                LIMIT 1
                """,
                params,
            )
            resume = cur.fetchone()

            if not resume:
                return None, None

            cur.execute(
                """
                SELECT *
                FROM job_web_resume_profiles
                WHERE resume_id = %s
                """,
                (resume["id"],),
            )
            profile = cur.fetchone()

    return dict(resume), dict(profile) if profile else None


def process_uploaded_resume(*, resume_id: int, storage_path: Path) -> tuple[dict[str, Any], dict[str, Any] | None]:
    raw_text = extract_text_from_docx(storage_path)
    update_resume_text(resume_id=resume_id, raw_text=raw_text)

    profile, metadata = normalize_resume_with_llm(raw_text)
    upsert_resume_profile(resume_id=resume_id, profile=profile, metadata=metadata)

    resume, saved_profile = fetch_resume_and_profile(resume_id=resume_id)
    return serialize_resume(resume), serialize_resume_profile(saved_profile)


@app.get("/healthz")
def healthz():
    try:
        ensure_runtime_state()
        return jsonify({"ok": True})
    except Exception as exc:
        logger.exception("Health check failed")
        return jsonify({"ok": False, "message": str(exc)}), 500


@app.get("/api/jobs")
def list_jobs():
    page, page_size = get_page_args()
    companies = get_companies_arg()

    try:
        return jsonify(fetch_jobs_from_postgres(page=page, page_size=page_size, companies=companies))
    except Exception as exc:
        logger.exception("Failed to load jobs from PostgreSQL")
        return jsonify({"error": "failed_to_load_jobs", "message": str(exc)}), 500


@app.get("/api/jobs/search")
def search_jobs():
    keyword = request.args.get("q", "").strip()
    page, page_size = get_page_args()
    companies = get_companies_arg()

    if not keyword:
        return jsonify(fetch_jobs_from_postgres(page=page, page_size=page_size, companies=companies))

    try:
        return jsonify(search_jobs_in_elasticsearch(keyword=keyword, page=page, page_size=page_size, companies=companies))
    except Exception as exc:
        logger.exception("Failed to search jobs in Elasticsearch")
        return jsonify({"error": "failed_to_search_jobs", "message": str(exc)}), 500


@app.get("/api/jobs/<job_id>")
def job_detail(job_id: str):
    try:
        with db_connect() as conn:
            with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
                cur.execute(
                    """
                    SELECT
                        job_id,
                        company,
                        title,
                        location,
                        locations,
                        level_guess,
                        team,
                        raw_description,
                        team_description,
                        responsibilities,
                        minimum_qualifications,
                        preferred_qualifications,
                        skills,
                        domains,
                        url,
                        posted_at,
                        first_seen_at,
                        last_seen_at,
                        is_active,
                        detail_extracted_at
                    FROM jobs
                    WHERE job_id = %s
                    """,
                    (job_id,),
                )
                row = cur.fetchone()

        if not row:
            return jsonify({"error": "job_not_found"}), 404

        return jsonify({"job": serialize_job(dict(row))})
    except Exception as exc:
        logger.exception("Failed to load job detail")
        return jsonify({"error": "failed_to_load_job_detail", "message": str(exc)}), 500


@app.post("/api/resumes")
def upload_resume():
    ensure_runtime_state()

    file_storage = request.files.get("file")
    if file_storage is None or not file_storage.filename:
        return jsonify({"error": "missing_file", "message": "A DOCX file is required."}), 400

    if not allowed_resume_file(file_storage.filename):
        return jsonify({"error": "invalid_file_type", "message": "Only .docx files are supported."}), 400

    storage_path, file_size = save_uploaded_resume(file_storage)
    resume_id = insert_resume_record(
        filename=file_storage.filename,
        content_type=file_storage.content_type,
        storage_path=storage_path,
        file_size_bytes=file_size,
    )

    try:
        resume, profile = process_uploaded_resume(resume_id=resume_id, storage_path=storage_path)
        return jsonify({"resume": resume, "profile": profile}), 201
    except Exception as exc:
        logger.exception("Failed to process uploaded resume")
        update_resume_failure(resume_id=resume_id, message=str(exc))
        resume, profile = fetch_resume_and_profile(resume_id=resume_id)
        return jsonify({
            "error": "resume_processing_failed",
            "message": str(exc),
            "resume": serialize_resume(resume),
            "profile": serialize_resume_profile(profile),
        }), 500


@app.get("/api/resumes/latest")
def latest_resume():
    ensure_runtime_state()
    resume, profile = fetch_resume_and_profile(latest=True)
    if not resume:
        return jsonify({"error": "resume_not_found"}), 404
    return jsonify({"resume": serialize_resume(resume), "profile": serialize_resume_profile(profile)})


@app.get("/api/resumes/<int:resume_id>")
def get_resume(resume_id: int):
    ensure_runtime_state()
    resume, profile = fetch_resume_and_profile(resume_id=resume_id)
    if not resume:
        return jsonify({"error": "resume_not_found"}), 404
    return jsonify({"resume": serialize_resume(resume), "profile": serialize_resume_profile(profile)})


@app.get("/api/resumes/<int:resume_id>/profile")
def get_resume_profile(resume_id: int):
    ensure_runtime_state()
    resume, profile = fetch_resume_and_profile(resume_id=resume_id)
    if not resume:
        return jsonify({"error": "resume_not_found"}), 404
    return jsonify({"resume": serialize_resume(resume), "profile": serialize_resume_profile(profile)})


@app.delete("/api/resumes/<int:resume_id>")
def delete_resume(resume_id: int):
    ensure_runtime_state()
    resume, _profile = fetch_resume_and_profile(resume_id=resume_id)
    if not resume:
        return jsonify({"error": "resume_not_found"}), 404

    storage_path = Path(resume["storage_path"])
    with db_connect() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM job_web_resumes WHERE id = %s", (resume_id,))
        conn.commit()

    try:
        if storage_path.exists():
            storage_path.unlink()
    except OSError:
        logger.warning("Failed to delete resume file %s", storage_path)

    return jsonify({"deleted": True, "resume_id": resume_id})


@app.route("/", defaults={"path": ""})
@app.route("/<path:path>")
def frontend(path: str):
    if path.startswith("api/") or path == "healthz":
        return jsonify({"error": "not_found"}), 404

    candidate = DIST_DIR / path
    if path and candidate.exists() and candidate.is_file():
        return send_from_directory(DIST_DIR, path)
    return send_from_directory(DIST_DIR, "index.html")


ensure_runtime_state()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8000)
